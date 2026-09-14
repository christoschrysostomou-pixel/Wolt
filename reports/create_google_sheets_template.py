#!/usr/bin/env python3
"""Create a Google Sheets-compatible workbook for Metro Cyprus report inputs."""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from generate_metro_cyprus_weekly_report import (
    ADDITION_DEDUCTION_METRICS,
    MONTHLY_PICKER_COLUMNS,
    OPERATIONS_METRICS,
    PICKER_METRICS,
    PURCHASE_METRICS,
    QUALITY_METRICS,
    WOLT_PLUS_METRICS,
)


WEEKLY_TEMPLATE = Path("data/weekly_metrics_template.csv")
MONTHLY_PICKER_TEMPLATE = Path("data/monthly_picker_metrics_template.csv")
DEFAULT_OUTPUT = Path("data/metro_cyprus_report_google_sheets_template.xlsx")

HEADER_FILL = PatternFill("solid", fgColor="00AEEF")
SECTION_FILL = PatternFill("solid", fgColor="E6F7FF")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(color="102A43", bold=True, size=14)
BODY_FONT = Font(color="334E68", size=10)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--weekly-template", type=Path, default=WEEKLY_TEMPLATE)
    parser.add_argument("--monthly-picker-template", type=Path, default=MONTHLY_PICKER_TEMPLATE)
    return parser.parse_args()


def load_headers(path: Path) -> list[str]:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        return next(reader)


def create_workbook(
    output_path: Path,
    weekly_template: Path = WEEKLY_TEMPLATE,
    monthly_picker_template: Path = MONTHLY_PICKER_TEMPLATE,
) -> None:
    weekly_headers = load_headers(weekly_template)
    monthly_headers = load_headers(monthly_picker_template)

    workbook = Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    add_instructions_sheet(workbook)
    add_metric_catalog_sheet(workbook)
    add_input_sheet(
        workbook,
        title="weekly_metrics",
        headers=weekly_headers,
        comments={
            "metric": "Use a metric name from the metric_catalog tab.",
            "venue_name": "Leave blank for merchant-level rows; populate for venue-level rows.",
            "value": "Numeric values only.",
            "dimension": "Optional chart grouping, such as hour, day, reason, or fee_type.",
            "dimension_value": "Optional label for the dimension, such as Monday or 13:00.",
        },
    )
    add_input_sheet(
        workbook,
        title="monthly_picker_metrics",
        headers=monthly_headers,
        comments={
            "Customer Feedback": "Optional free-text customer feedback, where available.",
        },
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)

    # Reopen once so generation fails early if the workbook is not readable.
    load_workbook(output_path, read_only=True).close()


def add_instructions_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("instructions")
    rows = [
        ["Metro Cyprus report input workbook"],
        [""],
        ["How to use"],
        ["1. Upload this .xlsx file to Google Drive or open it with Google Sheets."],
        ["2. Populate weekly_metrics with last-week report rows."],
        ["3. Populate monthly_picker_metrics with picker detail rows."],
        ["4. Download each populated tab as CSV and place them in data/:"],
        ["   - weekly_metrics.csv"],
        ["   - monthly_picker_metrics.csv"],
        ["5. Run: python3 reports/generate_metro_cyprus_weekly_report.py"],
        [""],
        ["Notes"],
        ["- Customer Feedback is optional and may be left blank when unavailable."],
        ["- Leave venue_name blank for merchant-level totals."],
        ["- Use metric names exactly as shown in the metric_catalog tab."],
    ]
    for row in rows:
        sheet.append(row)

    sheet["A1"].font = TITLE_FONT
    for row in sheet.iter_rows(min_row=3, max_col=1):
        row[0].font = BODY_FONT
    sheet.column_dimensions["A"].width = 90


def add_metric_catalog_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("metric_catalog")
    sheet.append(["Section", "Metric"])
    style_header(sheet, 1, 2)

    sections = [
        ("Purchases", PURCHASE_METRICS),
        ("Wolt+", WOLT_PLUS_METRICS),
        ("Operations", OPERATIONS_METRICS),
        ("Quality", QUALITY_METRICS),
        ("Additions & Deductions", ADDITION_DEDUCTION_METRICS),
        ("Picker", PICKER_METRICS),
    ]
    for section, metrics in sections:
        for metric in metrics:
            sheet.append([section, metric])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    set_widths(sheet, {"A": 28, "B": 42})
    for row in sheet.iter_rows(min_row=2):
        row[0].fill = SECTION_FILL


def add_input_sheet(
    workbook: Workbook,
    title: str,
    headers: list[str],
    comments: dict[str, str],
) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    style_header(sheet, 1, len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for index, header in enumerate(headers, start=1):
        column_letter = get_column_letter(index)
        sheet.column_dimensions[column_letter].width = max(18, min(42, len(header) + 8))
        if header in comments:
            sheet.cell(row=1, column=index).comment = Comment(comments[header], "Cursor")

    for row_index in range(2, 102):
        for column_index in range(1, len(headers) + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    if "Customer Feedback" in headers:
        feedback_column = get_column_letter(headers.index("Customer Feedback") + 1)
        sheet.column_dimensions[feedback_column].width = 48


def style_header(sheet, row_index: int, column_count: int) -> None:
    for column_index in range(1, column_count + 1):
        cell = sheet.cell(row=row_index, column=column_index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def set_widths(sheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def main() -> None:
    args = parse_args()
    create_workbook(args.output, args.weekly_template, args.monthly_picker_template)
    print(args.output)


if __name__ == "__main__":
    main()
