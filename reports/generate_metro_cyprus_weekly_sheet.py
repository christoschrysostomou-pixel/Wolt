#!/usr/bin/env python3
"""Generate the Metro Cyprus weekly report as a downloadable spreadsheet.

Produces an .xlsx workbook that mirrors the PDF report sections (merchant
overview, per-venue summary, and every metric section) so it can be
downloaded and opened directly in Google Sheets (Drive upload, or
File > Import in Sheets) as well as Excel or LibreOffice.

It reads the same CSV exports as the PDF generator, via
reports/metro_cyprus_data.py, so the two outputs never drift.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_REPO_ROOT = Path(__file__).resolve().parent.parent
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from reports.metro_cyprus_data import (  # noqa: E402
    MERCHANT_NAME,
    MetricRecord,
    MONTHLY_PICKER_COLUMNS,
    PER_VENUE_SUMMARY_METRICS,
    REPORT_SECTIONS,
    aggregate_metric,
    known_venues,
    load_monthly_picker,
    load_weekly_metrics,
    output_path_for,
    parse_date,
    previous_full_week,
    records_for_metric,
    records_for_venue,
)

HEADER_FILL = PatternFill(start_color="00AEEF", end_color="00AEEF", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FONT = Font(size=14, bold=True, color="102A43")
SUBTITLE_FONT = Font(size=10, italic=True, color="52606D")
NOTE_FONT = Font(size=9, italic=True, color="829AB1")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--weekly-metrics",
        type=Path,
        default=Path("data/weekly_metrics.csv"),
        help="CSV with weekly metrics in long format.",
    )
    parser.add_argument(
        "--monthly-picker",
        type=Path,
        default=Path("data/monthly_picker_metrics.csv"),
        help="CSV with monthly picker detail rows.",
    )
    parser.add_argument(
        "--period-start",
        type=str,
        default=None,
        help="Report period start date, YYYY-MM-DD. Defaults to previous Monday.",
    )
    parser.add_argument(
        "--period-end",
        type=str,
        default=None,
        help="Report period end date, YYYY-MM-DD. Defaults to previous Sunday.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path. Defaults to reports/output/metro_cyprus_weekly_report_<period>.xlsx.",
    )
    return parser.parse_args()


def style_title(ws: Worksheet, title: str, subtitle: str) -> int:
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    return 4


def style_header_row(ws: Worksheet, row: int, columns: list[str]) -> None:
    for col_index, column in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=col_index, value=column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate


def autosize_columns(ws: Worksheet, columns: list[str], sample_rows: list[list[object]]) -> None:
    for col_index, column in enumerate(columns, start=1):
        max_length = len(str(column))
        for row in sample_rows:
            if col_index - 1 < len(row):
                max_length = max(max_length, len(str(row[col_index - 1])))
        ws.column_dimensions[get_column_letter(col_index)].width = min(max(max_length + 2, 10), 48)


def add_note(ws: Worksheet, row: int, note: str) -> None:
    cell = ws.cell(row=row, column=1, value=note)
    cell.font = NOTE_FONT


def build_overview_sheet(
    wb: Workbook,
    records: list[MetricRecord],
    start,
    end,
    data_note: str,
) -> None:
    ws = wb.active
    ws.title = "Merchant Overview"
    next_row = style_title(
        ws,
        f"{MERCHANT_NAME} Weekly Report",
        f"Merchant preview | {start:%d %b %Y} - {end:%d %b %Y}",
    )

    overview_metrics = [
        "Group orders",
        "Venue Group Gross Value",
        "W+ Orders",
        "W+ New Users",
        "Ops Performance",
        "POFR Group %",
        "Substitution %",
        "% Unfulfilled Items",
        "% Late Orders",
        "Lost sales %GOV",
        "Lost sales GOV EUR",
    ]
    columns = ["Metric", "Value", "Unit"]
    style_header_row(ws, next_row, columns)
    rows: list[list[object]] = []
    for metric in overview_metrics:
        value, unit = aggregate_metric(records, metric)
        rows.append([metric, value if value is not None else "N/A", unit])
    for offset, row in enumerate(rows, start=1):
        for col_index, cell_value in enumerate(row, start=1):
            ws.cell(row=next_row + offset, column=col_index, value=cell_value)
    autosize_columns(ws, columns, rows)
    add_note(ws, next_row + len(rows) + 2, data_note)


def build_venue_summary_sheet(
    wb: Workbook,
    records: list[MetricRecord],
    venues: list[str],
    start,
    end,
    data_note: str,
) -> None:
    ws = wb.create_sheet("Venue Summary")
    next_row = style_title(
        ws,
        "Per-Venue Summary",
        f"Venue performance | {start:%d %b %Y} - {end:%d %b %Y}",
    )

    columns = ["Venue Name", *PER_VENUE_SUMMARY_METRICS]
    style_header_row(ws, next_row, columns)

    rows: list[list[object]] = []
    if venues:
        for venue in venues:
            venue_records = records_for_venue(records, venue)
            row: list[object] = [venue]
            for metric in PER_VENUE_SUMMARY_METRICS:
                value, _unit = aggregate_metric(venue_records, metric)
                row.append(value if value is not None else "N/A")
            rows.append(row)
    else:
        rows.append(["No venue-level weekly metrics were available."] + ["N/A"] * len(PER_VENUE_SUMMARY_METRICS))

    for offset, row in enumerate(rows, start=1):
        for col_index, cell_value in enumerate(row, start=1):
            ws.cell(row=next_row + offset, column=col_index, value=cell_value)
    autosize_columns(ws, columns, rows)
    add_note(ws, next_row + len(rows) + 2, data_note)


def build_section_sheet(
    wb: Workbook,
    title: str,
    records: list[MetricRecord],
    metrics: list[str],
    start,
    end,
    data_note: str,
) -> None:
    ws = wb.create_sheet(title[:31])
    next_row = style_title(ws, title, f"{start:%d %b %Y} - {end:%d %b %Y}")

    columns = [
        "Metric",
        "Venue Name",
        "Value",
        "Unit",
        "Dimension",
        "Dimension Value",
        "Period Start",
        "Period End",
    ]
    style_header_row(ws, next_row, columns)

    rows: list[list[object]] = []
    for metric in metrics:
        metric_records = records_for_metric(records, metric)
        if not metric_records:
            rows.append([metric, "", "N/A", "", "", "", "", ""])
            continue
        for record in metric_records:
            rows.append(
                [
                    record.metric,
                    record.venue_name,
                    record.value,
                    record.unit,
                    record.dimension,
                    record.dimension_value,
                    record.period_start,
                    record.period_end,
                ]
            )

    for offset, row in enumerate(rows, start=1):
        for col_index, cell_value in enumerate(row, start=1):
            ws.cell(row=next_row + offset, column=col_index, value=cell_value)
    autosize_columns(ws, columns, rows)
    add_note(ws, next_row + len(rows) + 2, data_note)


def build_monthly_picker_sheet(
    wb: Workbook,
    rows: list[dict[str, str]],
    start,
    end,
    data_note: str,
) -> None:
    ws = wb.create_sheet("Monthly Picker Metrics")
    next_row = style_title(
        ws,
        "Monthly Picker Metrics",
        f"Detail export included with weekly report | {start:%d %b %Y} - {end:%d %b %Y}",
    )

    style_header_row(ws, next_row, MONTHLY_PICKER_COLUMNS)

    table_rows: list[list[object]] = [
        [row.get(column, "") for column in MONTHLY_PICKER_COLUMNS] for row in rows
    ]
    if not table_rows:
        table_rows = [["No monthly picker detail export was provided."] + [""] * (len(MONTHLY_PICKER_COLUMNS) - 1)]

    for offset, row in enumerate(table_rows, start=1):
        for col_index, cell_value in enumerate(row, start=1):
            ws.cell(row=next_row + offset, column=col_index, value=cell_value)
    autosize_columns(ws, MONTHLY_PICKER_COLUMNS, table_rows)
    add_note(ws, next_row + len(table_rows) + 2, data_note)


def generate_sheet(
    weekly_records: list[MetricRecord],
    monthly_picker_rows: list[dict[str, str]],
    start,
    end,
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    data_note = (
        "Data source: CSV exports supplied to the generator."
        if weekly_records or monthly_picker_rows
        else "Data source unavailable in this run; all N/A rows require weekly and monthly CSV exports."
    )

    wb = Workbook()
    build_overview_sheet(wb, weekly_records, start, end, data_note)
    venues = known_venues(weekly_records)
    build_venue_summary_sheet(wb, weekly_records, venues, start, end, data_note)
    for title, metrics in REPORT_SECTIONS:
        build_section_sheet(wb, title, weekly_records, metrics, start, end, data_note)
    build_monthly_picker_sheet(wb, monthly_picker_rows, start, end, data_note)

    wb.save(output_path)


def main() -> None:
    args = parse_args()
    default_start, default_end = previous_full_week()
    start = parse_date(args.period_start, default_start)
    end = parse_date(args.period_end, default_end)
    output_path = args.output or output_path_for(start, end, "xlsx")

    weekly_records = load_weekly_metrics(args.weekly_metrics)
    monthly_picker_rows = load_monthly_picker(args.monthly_picker)
    generate_sheet(weekly_records, monthly_picker_rows, start, end, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
