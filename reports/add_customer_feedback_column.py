#!/usr/bin/env python3
"""Add a Customer Feedback column to a Monthly Picker Metrics workbook.

Appends a "Customer Feedback" column to every worksheet that looks like a
Monthly Picker Metrics export (has "Purchase ID" and "Order Number" headers).
Values are only filled in when a feedback lookup CSV is supplied and a row
matches; unmatched rows are left blank rather than fabricated, since customer
feedback text is not something this tool should invent.

Typical feedback source: an export from Snowflake keyed by Purchase ID or
Order Number with a "Customer Feedback" column. Once that export is
available, pass it via --feedback-csv to merge it in.
"""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

FEEDBACK_COLUMN = "Customer Feedback"
KEY_COLUMNS = ("Purchase ID", "Order Number")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Source .xlsx workbook.")
    parser.add_argument("output", type=Path, help="Destination .xlsx workbook.")
    parser.add_argument(
        "--feedback-csv",
        type=Path,
        default=None,
        help=(
            "Optional CSV with a key column (Purchase ID or Order Number) and a "
            "Customer Feedback column to merge in. Rows without a match are left blank."
        ),
    )
    return parser.parse_args()


def load_feedback_lookup(path: Path | None) -> tuple[dict[str, str], str | None]:
    if path is None:
        return {}, None
    if not path.exists():
        raise FileNotFoundError(f"Feedback CSV not found: {path}")

    import csv

    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.DictReader(handle)
        fieldnames = reader.fieldnames or []
        key_column = next((column for column in KEY_COLUMNS if column in fieldnames), None)
        if key_column is None or FEEDBACK_COLUMN not in fieldnames:
            raise ValueError(
                f"{path} must contain a '{FEEDBACK_COLUMN}' column and one of {KEY_COLUMNS}"
            )
        lookup = {
            str(row[key_column]).strip(): str(row[FEEDBACK_COLUMN] or "").strip()
            for row in reader
            if row.get(key_column)
        }
    return lookup, key_column


def is_monthly_picker_sheet(headers: list[object]) -> bool:
    return "Purchase ID" in headers and "Order Number" in headers


def add_feedback_column(
    input_path: Path,
    output_path: Path,
    feedback_csv: Path | None = None,
) -> dict[str, dict[str, int]]:
    lookup, key_column = load_feedback_lookup(feedback_csv)
    workbook = load_workbook(input_path)
    summary: dict[str, dict[str, int]] = {}

    for sheet in workbook.worksheets:
        header_row = next(sheet.iter_rows(min_row=1, max_row=1))
        headers = [cell.value for cell in header_row]
        if not is_monthly_picker_sheet(headers):
            continue

        if FEEDBACK_COLUMN in headers:
            feedback_col_idx = headers.index(FEEDBACK_COLUMN) + 1
        else:
            feedback_col_idx = len(headers) + 1
            style_source = sheet.cell(row=1, column=len(headers))
            new_header = sheet.cell(row=1, column=feedback_col_idx, value=FEEDBACK_COLUMN)
            new_header.font = copy(style_source.font)
            new_header.fill = copy(style_source.fill)
            new_header.border = copy(style_source.border)
            new_header.alignment = copy(style_source.alignment)
            new_header.number_format = "@"
            sheet.column_dimensions[new_header.column_letter].width = 42

        purchase_col = headers.index("Purchase ID") + 1
        order_col = headers.index("Order Number") + 1

        total_rows = 0
        matched_rows = 0
        for row_idx in range(2, sheet.max_row + 1):
            purchase_id = sheet.cell(row=row_idx, column=purchase_col).value
            order_number = sheet.cell(row=row_idx, column=order_col).value
            if purchase_id is None and order_number is None:
                continue
            total_rows += 1

            feedback_cell = sheet.cell(row=row_idx, column=feedback_col_idx)
            if lookup:
                key_value = purchase_id if key_column == "Purchase ID" else order_number
                feedback = lookup.get(str(key_value).strip(), "") if key_value is not None else ""
                if feedback:
                    matched_rows += 1
                    feedback_cell.value = feedback
                elif feedback_cell.value is None:
                    feedback_cell.value = ""
            elif feedback_cell.value is None:
                feedback_cell.value = ""

        summary[sheet.title] = {"rows": total_rows, "matched": matched_rows}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return summary


def main() -> None:
    args = parse_args()
    summary = add_feedback_column(args.input, args.output, args.feedback_csv)
    if not summary:
        print("No Monthly Picker Metrics sheets found (need 'Purchase ID' and 'Order Number' headers).")
    for sheet_name, stats in summary.items():
        print(f"{sheet_name}: {stats['rows']} rows, {stats['matched']} with feedback")
    print(args.output)


if __name__ == "__main__":
    main()
