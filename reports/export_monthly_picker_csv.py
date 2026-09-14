#!/usr/bin/env python3
"""Export a Monthly Picker Metrics workbook to the report generator's CSV schema.

Reads every worksheet that looks like a Monthly Picker Metrics export (has
"Purchase ID" and "Order Number" headers) and concatenates them into a single
CSV matching data/monthly_picker_metrics_template.csv, so it can be consumed
directly by reports/generate_metro_cyprus_weekly_report.py and
reports/generate_metro_cyprus_weekly_sheet.py.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import sys
from pathlib import Path

from openpyxl import load_workbook

_REPO_ROOT = Path(__file__).resolve().parent.parent
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from reports.metro_cyprus_data import MONTHLY_PICKER_COLUMNS  # noqa: E402


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Source .xlsx workbook (one or more monthly sheets).")
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("data/monthly_picker_metrics.csv"),
        help="Destination CSV, defaults to data/monthly_picker_metrics.csv.",
    )
    return parser.parse_args()


def format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.strftime("%Y-%m-%d")
        return value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value)


def export(input_path: Path, output_path: Path) -> int:
    workbook = load_workbook(input_path, data_only=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    row_count = 0
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(MONTHLY_PICKER_COLUMNS)

        for sheet in workbook.worksheets:
            header_row = next(sheet.iter_rows(min_row=1, max_row=1))
            headers = [cell.value for cell in header_row]
            if "Purchase ID" not in headers or "Order Number" not in headers:
                continue

            column_index = {name: idx for idx, name in enumerate(headers) if name is not None}
            purchase_idx = column_index["Purchase ID"]

            for row in sheet.iter_rows(min_row=2, values_only=True):
                if purchase_idx >= len(row) or row[purchase_idx] is None:
                    continue
                out_row = [
                    format_cell(row[column_index[column]]) if column in column_index and column_index[column] < len(row) else ""
                    for column in MONTHLY_PICKER_COLUMNS
                ]
                writer.writerow(out_row)
                row_count += 1

    return row_count


def main() -> None:
    args = parse_args()
    count = export(args.input, args.output)
    print(f"Wrote {count} rows to {args.output}")


if __name__ == "__main__":
    main()
