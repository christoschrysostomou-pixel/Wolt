#!/usr/bin/env python3
"""
Fill missing weeks in the "CYP Metro Sales Monitoring.xlsx" workbook.

The workbook has two sheets:
  - "Data": the raw weekly numbers pulled from Looker (explore
    `core.f_purchases`, grouped by "Purchases Time Delivered Week of Year",
    split by "Purchases Time Delivered Year"). This is the only sheet that
    needs to be edited by hand/by this script.
  - "Analysis": pulls everything from "Data" via XLOOKUP formulas, so it
    recalculates automatically in Excel/Sheets once "Data" is updated. It
    does not need to be touched directly.

"Data" sheet layout (row 1 = year, row 2 = column labels, data from row 3):
    A: week number (duplicate of B, kept for backward compat)
    B: Purchases Time Delivered Week of Year
    C: Purchases Number of Orders            (current year)
    D: Purchases GOV Total Euro              (current year)
    E: % Change  = (C - G) / G               (YoY orders change)
    F: % GOV     = (D - H) / H               (YoY GOV change)
    G: Purchases Number of Orders            (prior year)
    H: Purchases GOV Total Euro              (prior year)
    I: % Change for the prior year column (always 0, kept for symmetry)
    J: % GOV for the prior year column (always 0, kept for symmetry)

This script only fills in C/D (and recomputes E/F) for the weeks you supply;
it never overwrites weeks that aren't included in the input, and it leaves
the "Analysis" sheet formulas untouched.

Usage:
    python update_data_sheet.py \
        --workbook "CYP Metro Sales Monitoring.xlsx" \
        --year 2025 \
        --input new_weeks.csv \
        --output "CYP Metro Sales Monitoring (updated).xlsx"

Where new_weeks.csv has the columns: week,orders,gov
(one row per ISO week you pulled from Looker, week = 1-53).

You can get `orders`/`gov` for the missing weeks straight out of the same
Looker explore that produced the existing rows:
https://looker.wolt.com/explore/core/f_purchases

Group by "Purchases Time Delivered Week of Year", filter to the market/year
used elsewhere in this workbook (see the existing "Data" sheet filters for
reference), and pull "Purchases Number of Orders" and
"Purchases GOV (Gross Order Value) Total Euro" for each missing week up to
(and including) the last fully completed week.
"""

import argparse
import csv
import sys

import openpyxl

WEEK_COL = "B"
CUR_ORDERS_COL = "C"
CUR_GOV_COL = "D"
CHANGE_COL = "E"
GOV_CHANGE_COL = "F"
PRIOR_ORDERS_COL = "G"
PRIOR_GOV_COL = "H"

FIRST_DATA_ROW = 3


def load_new_weeks(path):
    """Read a CSV with columns week,orders,gov into {week: (orders, gov)}."""
    weeks = {}
    with open(path, newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            week = int(float(row["week"]))
            orders = float(row["orders"])
            gov = float(row["gov"])
            weeks[week] = (orders, gov)
    return weeks


def update_workbook(workbook_path, year, new_weeks, output_path):
    wb = openpyxl.load_workbook(workbook_path, data_only=False)
    ws = wb["Data"]

    # Sanity check the sheet is laid out the way this script expects.
    header_year = ws["C1"].value
    if header_year not in (year, float(year)):
        print(
            f"warning: Data!C1 is {header_year!r}, expected {year}. "
            "Double check the column layout before trusting the output.",
            file=sys.stderr,
        )

    updated = []
    row = FIRST_DATA_ROW
    while ws[f"{WEEK_COL}{row}"].value is not None:
        week_val = ws[f"{WEEK_COL}{row}"].value
        week = int(float(week_val))
        if week in new_weeks:
            orders, gov = new_weeks[week]
            prior_orders = ws[f"{PRIOR_ORDERS_COL}{row}"].value or 0
            prior_gov = ws[f"{PRIOR_GOV_COL}{row}"].value or 0

            ws[f"{CUR_ORDERS_COL}{row}"] = orders
            ws[f"{CUR_GOV_COL}{row}"] = gov
            ws[f"{CHANGE_COL}{row}"] = (
                (orders - prior_orders) / prior_orders if prior_orders else 0
            )
            ws[f"{GOV_CHANGE_COL}{row}"] = (
                (gov - prior_gov) / prior_gov if prior_gov else 0
            )
            updated.append(week)
        row += 1

    missing = sorted(set(new_weeks) - set(updated))
    if missing:
        print(
            f"warning: weeks {missing} were not found as existing rows in "
            "the Data sheet and were skipped. Add rows for them first.",
            file=sys.stderr,
        )

    wb.save(output_path)
    print(f"Updated weeks: {sorted(updated)}")
    print(f"Saved to: {output_path}")
    print(
        "Open the file in Excel and let it recalculate (or press "
        "Ctrl+Alt+F9) so the 'Analysis' sheet formulas pick up the new "
        "values."
    )


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, help="Path to the source .xlsx")
    parser.add_argument(
        "--year", required=True, type=int, help="Year the new weeks belong to"
    )
    parser.add_argument(
        "--input",
        required=True,
        help="CSV file with columns: week,orders,gov",
    )
    parser.add_argument(
        "--output", required=True, help="Path to write the updated .xlsx to"
    )
    args = parser.parse_args()

    new_weeks = load_new_weeks(args.input)
    update_workbook(args.workbook, args.year, new_weeks, args.output)


if __name__ == "__main__":
    main()
